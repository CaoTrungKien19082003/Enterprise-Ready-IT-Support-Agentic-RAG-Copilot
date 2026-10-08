from pathlib import Path
from typing import Iterable 
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader 
from docx import Document as DocxDocument


SUPPORTED = [".pdf", ".txt", ".md", ".docx"] 
#if you want to support more file types, you can add them here and community loaders for them. 
# For example, for .csv files, you can use the CSVLoader from langchain_community.document_loaders

def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    elif suffix == ".txt" or suffix == ".md":
        return TextLoader(str(path), encoding="utf-8").load()
    elif suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    else:
        raise ValueError(f"Unsupported file type: {path.suffix}")


def chunk_documents(documents: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(documents)